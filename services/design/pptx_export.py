"""Deck → PowerPoint (python-pptx, imported lazily).

Two modes (`options.mode`):

* ``screenshots`` (default) — one full-bleed picture per slide, speaker notes attached.
  Pixel-exact, not editable.
* ``editable`` — per slide, the text is read from the DOM (rects, runs, computed font /
  size / weight / colour / alignment / line height), the page is re-captured with every
  extracted glyph made transparent (``-webkit-text-fill-color`` — ``color`` itself is left
  alone so borders and icons drawn in ``currentColor`` survive), and the picture becomes
  the slide background under real, editable text boxes.

Options: ``fontSwaps`` ``{"Inter": "Arial"}`` (first family of each stack, case-insensitive),
``hideSelectors`` (hidden in both modes; deck controls are always hidden), ``scale`` (1–2,
screenshot density), and:

* ``googleFontImports`` — families (``"Inter:wght@400;700"``) or full ``fonts.googleapis.com``
  URLs, loaded into the page before capture (for a font the page names but never imports);
* ``resetTransformSelector`` — CSS selector whose ``transform`` is forced to ``none`` before
  capture (a scaled / letterboxed stage that would otherwise shrink every slide);
* ``slides`` — per-slide overrides ``[{index?, selector?, showJs?, delay?}]``. When given it
  defines the slide list: entry *i* goes to deck slide ``index`` (default *i*, decks only),
  runs ``showJs`` (awaited JS in the page — open a tab, advance a stepper), waits ``delay``
  ms (≤ 10 000), then captures ``selector``'s box (default: the whole slide / viewport). On
  a page that is not a deck this is how several slides come out of one page;
* ``save_to_project_path`` — also write the .pptx into the project (``exports/deck.pptx``),
  validated like any project write and recorded as a version.

Validation flags on the job: ``duplicate_adjacent`` (two consecutive slides render
identically — usually navigation that did not move), ``slide_size_mismatch`` (a slide's
box differs from the declared canvas, or the canvas is not 16:9), ``no_speaker_notes``
(missing, or not one entry per slide), plus ``web_fonts`` (families PowerPoint will
substitute unless swapped) and ``deck_nav_forced``.

Notes come from ``<script type="application/json" id="td-speaker-notes">`` (or
``id="speaker-notes"``): a JSON array, one entry per slide.
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import logging
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from services.design import render

log = logging.getLogger("telecode.services.design.pptx_export")

SLIDE_WIDTH_EMU = 12192000  # 13.333 in — PowerPoint's 16:9 default
_OFFICE_FONTS = {
    "arial", "calibri", "cambria", "candara", "consolas", "constantia", "corbel", "courier new", "georgia",
    "segoe ui", "tahoma", "times new roman", "trebuchet ms", "verdana", "aptos", "helvetica", "garamond",
    "century gothic", "franklin gothic", "gill sans mt", "lucida console", "palatino linotype", "book antiqua",
    "impact", "segoe ui semibold", "segoe ui light", "arial black", "sans-serif", "serif", "monospace",
}
_GENERIC = {"sans-serif": "Arial", "serif": "Times New Roman", "monospace": "Consolas", "system-ui": "Segoe UI",
            "ui-sans-serif": "Segoe UI", "ui-serif": "Georgia", "ui-monospace": "Consolas",
            "-apple-system": "Segoe UI", "blinkmacsystemfont": "Segoe UI"}

_TEXT_JS = r"""
(slideIndex, W, selector) => {
  const root = selector ? document.querySelector(selector)
      : slideIndex ? window.__tdDeck.slides()[slideIndex - 1] : document.body;
  if (!root) return {factor: 1, blocks: []};
  const rr = selector || slideIndex ? root.getBoundingClientRect() : {x: 0, y: 0, width: innerWidth, height: innerHeight};
  const factor = rr.width ? W / rr.width : 1;
  const SKIP = new Set(['SCRIPT', 'STYLE', 'NOSCRIPT', 'TEMPLATE', 'CANVAS', 'IFRAME', 'VIDEO', 'AUDIO',
                        'TEXTAREA', 'INPUT', 'SELECT', 'OPTION', 'IMG', 'PICTURE', 'OBJECT', 'EMBED']);
  const cv = document.createElement('canvas'); cv.width = cv.height = 1;
  const cx = cv.getContext('2d', {willReadFrequently: true});
  const cache = new Map();
  const rgba = (c) => {
    if (cache.has(c)) return cache.get(c);
    cx.clearRect(0, 0, 1, 1); cx.fillStyle = '#000'; cx.fillStyle = c; cx.fillRect(0, 0, 1, 1);
    const d = cx.getImageData(0, 0, 1, 1).data; const v = [d[0], d[1], d[2], +(d[3] / 255).toFixed(3)];
    cache.set(c, v); return v;
  };
  const blockish = (cs) => cs.display !== 'inline' && cs.display !== 'contents';
  const hidden = (cs) => cs.display === 'none' || cs.visibility === 'hidden' || +cs.opacity < 0.02;
  const tt = (t, m) => m === 'uppercase' ? t.toUpperCase() : m === 'lowercase' ? t.toLowerCase()
      : m === 'capitalize' ? t.replace(/(^|\s)(\p{L})/gu, (a, s, c) => s + c.toUpperCase()) : t;
  const style = (cs) => {
    const fill = cs.webkitTextFillColor && cs.webkitTextFillColor !== cs.color ? cs.webkitTextFillColor : cs.color;
    let color = rgba(fill);
    if (color[3] === 0 && cs.backgroundClip === 'text') color = rgba(cs.color === fill ? '#000' : cs.color);
    return {font: cs.fontFamily, size: parseFloat(cs.fontSize), weight: parseInt(cs.fontWeight) || 400,
            italic: cs.fontStyle === 'italic' || cs.fontStyle.startsWith('oblique'),
            underline: (cs.textDecorationLine || '').includes('underline'),
            strike: (cs.textDecorationLine || '').includes('line-through'),
            color, ls: cs.letterSpacing === 'normal' ? 0 : parseFloat(cs.letterSpacing) || 0};
  };
  const pseudo = (el, which) => {
    const c = getComputedStyle(el, which).content;
    if (!c || c === 'none' || c === 'normal') return '';
    const m = c.match(/^"([\s\S]*)"$/); return m ? m[1] : '';
  };
  const collect = (block) => {
    const runs = [], rects = [];
    const walk = (node) => {
      for (const ch of node.childNodes) {
        if (ch.nodeType === 3) {
          const raw = ch.textContent; if (!raw) continue;
          const cs = getComputedStyle(ch.parentElement);
          const pre = /^pre/.test(cs.whiteSpace);
          const t = pre ? raw : raw.replace(/\s+/g, ' ');
          if (!t.trim() && !runs.length) continue;
          const rg = document.createRange(); rg.selectNodeContents(ch);
          for (const r of rg.getClientRects()) if (r.width > 0.5 && r.height > 0.5) rects.push(r);
          runs.push({t: tt(t, cs.textTransform), pre, ...style(cs)});
        } else if (ch.nodeType === 1) {
          if (ch instanceof SVGElement || SKIP.has(ch.tagName)) continue;
          if (ch.tagName === 'BR') { runs.push({br: true}); continue; }
          const cs = getComputedStyle(ch);
          if (hidden(cs) || blockish(cs)) continue;
          const b = pseudo(ch, '::before'); if (b) runs.push({t: b, ...style(cs)});
          walk(ch);
          const a = pseudo(ch, '::after'); if (a) runs.push({t: a, ...style(cs)});
        }
      }
    };
    const bcs = getComputedStyle(block);
    if (bcs.display === 'list-item' && bcs.listStyleType !== 'none') {
      const idx = [...block.parentElement.children].filter(c => c.tagName === 'LI').indexOf(block) + 1;
      runs.push({t: /decimal|numer/.test(bcs.listStyleType) ? `${idx}. ` : '• ', ...style(bcs)});
    }
    const b = pseudo(block, '::before'); if (b) runs.push({t: b, ...style(bcs)});
    walk(block);
    const a = pseudo(block, '::after'); if (a) runs.push({t: a, ...style(bcs)});
    return {runs, rects};
  };
  const blocks = [];
  const visit = (el) => {
    if (el instanceof SVGElement || SKIP.has(el.tagName)) return;
    if (el.closest('.deck-controls, [data-td-rt]')) return;
    const cs = getComputedStyle(el);
    if (cs.display === 'none') return;
    if ((el === root || blockish(cs)) && cs.visibility !== 'hidden' && +cs.opacity >= 0.02) {
      const {runs, rects} = collect(el);
      if (rects.length && runs.some(r => r.t && r.t.trim())) {
        let x0 = Infinity, y0 = Infinity, x1 = -Infinity, y1 = -Infinity;
        const tops = [];
        for (const r of rects) {
          x0 = Math.min(x0, r.left); y0 = Math.min(y0, r.top); x1 = Math.max(x1, r.right); y1 = Math.max(y1, r.bottom);
          if (!tops.some(t => Math.abs(t - r.top) < 3)) tops.push(r.top);
        }
        const size = parseFloat(cs.fontSize);
        const lh = cs.lineHeight === 'normal' ? size * 1.2 : parseFloat(cs.lineHeight);
        let align = cs.textAlign;
        if (align === 'start') align = cs.direction === 'rtl' ? 'right' : 'left';
        if (align === 'end') align = cs.direction === 'rtl' ? 'left' : 'right';
        blocks.push({x: (x0 - rr.x) * factor, y: (y0 - rr.y) * factor, w: (x1 - x0) * factor, h: (y1 - y0) * factor,
                     lines: tops.length, align, line_height: lh, runs, opacity: +cs.opacity});
        el.setAttribute('data-td-rt-txt', '');
      }
    }
    for (const c of el.children) visit(c);
  };
  visit(root);
  return {factor, blocks};
}
"""

_HIDE_TEXT_CSS = (
    "[data-td-rt-txt], [data-td-rt-txt] * { -webkit-text-fill-color: transparent !important;"
    " text-shadow: none !important; text-decoration-color: transparent !important;"
    " -webkit-text-stroke: 0 !important; }"
    "[data-td-rt-txt]::marker, [data-td-rt-txt] *::marker { color: transparent !important; }"
)


def _first_family(stack: str) -> str:
    return (stack.split(",")[0] if stack else "").strip().strip("'\"")


def _map_font(stack: str, swaps: Dict[str, str]) -> str:
    fam = _first_family(stack)
    low = fam.lower()
    for k, v in swaps.items():
        if k.lower() == low:
            return v
    return _GENERIC.get(low, fam) or "Arial"


MAX_SLIDE_SPECS = 200
_FONTS_HOST = "https://fonts.googleapis.com/"
_FAMILY_RE = re.compile(r"^[A-Za-z0-9 ]{1,60}(?::[A-Za-z0-9@;,.]{1,120})?$")
_URL_BAD = set("\\\"'<>() ")


def _clean_selector(sel: Any) -> Optional[str]:
    if not isinstance(sel, str):
        return None
    sel = sel.strip()
    if not sel or len(sel) > 300 or "{" in sel or "}" in sel or "<" in sel:
        return None
    return sel


def font_import_urls(items: Any) -> List[str]:
    """`googleFontImports` → stylesheet URLs on fonts.googleapis.com only (the preview
    CSP allows no other stylesheet host, and a caller-chosen host would be an SSRF
    through the headless browser)."""
    urls: List[str] = []
    for it in items if isinstance(items, list) else []:
        if not isinstance(it, str) or not it.strip():
            continue
        it = it.strip()
        if it.startswith(_FONTS_HOST) and not any(c in _URL_BAD for c in it):
            urls.append(it)
        elif _FAMILY_RE.match(it):
            fam, _, axes = it.partition(":")
            q = fam.strip().replace(" ", "+") + (":" + axes if axes else "")
            urls.append(f"{_FONTS_HOST}css2?family={q}&display=swap")
        if len(urls) >= 20:
            break
    return urls


def slide_specs(opts: Dict[str, Any]) -> Optional[List[Dict[str, Any]]]:
    """`slides` option → [{index, selector, showJs, delay}] (None when not given)."""
    raw = opts.get("slides")
    if not isinstance(raw, list) or not raw:
        return None
    out: List[Dict[str, Any]] = []
    for sp in raw[:MAX_SLIDE_SPECS]:
        sp = sp if isinstance(sp, dict) else {}
        try:
            idx = int(sp["index"]) if sp.get("index") is not None else None
        except (TypeError, ValueError):
            idx = None
        try:
            delay = max(0, min(10000, int(sp.get("delay") or 0)))
        except (TypeError, ValueError):
            delay = 0
        js = sp.get("showJs")
        out.append({"index": idx if idx and idx > 0 else None, "selector": _clean_selector(sp.get("selector")),
                    "showJs": js[:20000] if isinstance(js, str) and js.strip() else None, "delay": delay})
    return out


async def _prepare_page(page, opts: Dict[str, Any]) -> None:
    urls = font_import_urls(opts.get("googleFontImports"))
    if urls:
        await page.add_style("".join(f"@import url('{u}');" for u in urls), "__td_rt_fonts")
        try:
            await page.eval("Promise.race([new Promise(r => setTimeout(r, 150)).then(() => document.fonts.ready)"
                            ".then(() => new Promise(r => setTimeout(r, 250))), new Promise(r => setTimeout(r, 4000))])",
                            timeout=10)
        except Exception as exc:
            log.info("pptx: font imports did not settle: %s", exc)
    reset = _clean_selector(opts.get("resetTransformSelector"))
    if reset:
        await page.add_style(f"{reset}{{transform:none !important}}", "__td_rt_reset")
        await page.frame()


async def _clip_of(page, selector: str) -> Optional[Dict[str, float]]:
    r = await page.call(
        "(sel) => { const el = document.querySelector(sel); if (!el) return null;"
        " const r = el.getBoundingClientRect(); return [r.x + scrollX, r.y + scrollY, r.width, r.height]; }",
        selector)
    if not r or r[2] < 1 or r[3] < 1:
        return None
    return {"x": r[0], "y": r[1], "width": r[2], "height": r[3]}


async def _capture_slides(ctx, mode: str) -> Dict[str, Any]:
    opts = ctx.options
    scale = max(1.0, min(2.0, float(opts.get("scale") or 1)))
    hide = [".deck-controls"] + [s for s in (opts.get("hideSelectors") or []) if isinstance(s, str)]
    specs = slide_specs(opts)
    slides: List[Dict[str, Any]] = []
    missing: List[int] = []
    js_failed: List[int] = []
    async with render.open_page(1920, 1080, scale) as page:
        info = await render.load_for_render(page, ctx.url())
        await page.hide(hide)
        await _prepare_page(page, opts)
        is_deck = bool(info.get("is_deck"))
        W, H = (int(info["w"]), int(info["h"])) if is_deck else (1920, 1080)
        if specs is None:
            specs = [{"index": i, "selector": None, "showJs": None, "delay": 0}
                     for i in range(1, (int(info["count"]) if is_deck else 1) + 1)]
        elif not is_deck and specs[0]["selector"]:
            # Slides cut out of one page: the slide size is the first region's.
            first = await _clip_of(page, specs[0]["selector"])
            if first:
                W, H = max(1, round(first["width"])), max(1, round(first["height"]))
        count = len(specs)
        for i, sp in enumerate(specs, start=1):
            deck_n = (sp["index"] or i) if is_deck else 0
            nav = await render.deck_go(page, deck_n) if deck_n else {}
            if sp["showJs"]:
                try:
                    await page.eval("(async () => {\n" + sp["showJs"] + "\n})()", timeout=30)
                except Exception as exc:
                    js_failed.append(i)
                    log.info("pptx: showJs on slide %d failed: %s", i, exc)
            if sp["delay"]:
                await asyncio.sleep(sp["delay"] / 1000.0)
            await page.frame()
            clip = await _clip_of(page, sp["selector"]) if sp["selector"] else None
            if sp["selector"] and not clip:
                missing.append(i)
            text = None
            if mode == "editable":
                text = await page.call(_TEXT_JS, deck_n, W, sp["selector"] if clip else None, timeout=60)
                await page.add_style(_HIDE_TEXT_CSS, "__td_rt_hidetext")
                await page.frame()
            png = await page.capture("png", clip=clip)
            if mode == "editable":
                await page.remove_style("__td_rt_hidetext")
                await page.eval("document.querySelectorAll('[data-td-rt-txt]').forEach(e => "
                                "e.removeAttribute('data-td-rt-txt'))", await_promise=False)
            slides.append({"png": png, "text": text, "forced": bool(nav.get("forced")),
                           "size": nav.get("size") if not clip else None,
                           "rect": [clip["width"], clip["height"]] if clip else None, "deck_n": deck_n})
            ctx.progress(0.05 + 0.75 * i / count, f"Slide {i}/{count}")
        errors = page.errors
    return {"info": info, "is_deck": is_deck, "W": W, "H": H, "slides": slides, "errors": errors,
            "missing_selectors": missing, "js_failed": js_failed}


def _build(cap: Dict[str, Any], mode: str, swaps: Dict[str, str], out: Path) -> Dict[str, Any]:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Emu, Pt

    W, H = cap["W"], cap["H"]
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_WIDTH_EMU)
    prs.slide_height = Emu(round(SLIDE_WIDTH_EMU * H / W))
    emu_px = SLIDE_WIDTH_EMU / W
    pt_px = emu_px / 12700.0
    notes = cap["info"].get("notes") if cap["is_deck"] else None
    blank = prs.slide_layouts[6]
    align_map = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
                 "justify": PP_ALIGN.JUSTIFY}
    fonts_used: Dict[str, str] = {}
    text_boxes = 0

    for idx, s in enumerate(cap["slides"]):
        slide = prs.slides.add_slide(blank)
        if s.get("rect"):
            # A selector region: full width, height from its own aspect (top-aligned), so
            # the text boxes — mapped with the same width factor — land on the picture.
            rw, rh = s["rect"]
            ph = min(int(prs.slide_height), round(int(prs.slide_width) * rh / max(1.0, rw)))
            slide.shapes.add_picture(io.BytesIO(s["png"]), 0, 0, width=prs.slide_width, height=Emu(ph))
        else:
            slide.shapes.add_picture(io.BytesIO(s["png"]), 0, 0, width=prs.slide_width, height=prs.slide_height)
        if mode == "editable" and s["text"]:
            for b in s["text"]["blocks"]:
                runs = b["runs"]
                # CSS whitespace collapsing across run boundaries, then trim the ends.
                prev_space = True
                for r in runs:
                    if r.get("br") or r.get("pre"):
                        prev_space = bool(r.get("br"))
                        continue
                    t = r["t"]
                    if prev_space:
                        t = t.lstrip(" ")
                    prev_space = t.endswith(" ")
                    r["t"] = t
                for r in reversed(runs):
                    if r.get("br"):
                        continue
                    r["t"] = r["t"].rstrip(" ") if not r.get("pre") else r["t"]
                    if r["t"]:
                        break
                if not any((r.get("t") or "").strip() for r in runs):
                    continue
                w = b["w"] * 1.06 + 2
                x = b["x"]
                if b["align"] == "center":
                    x -= (w - b["w"]) / 2
                elif b["align"] == "right":
                    x -= (w - b["w"])
                x, y = max(0.0, x), max(0.0, b["y"])
                w = min(w, W - x) if W - x > 4 else w
                tb = slide.shapes.add_textbox(Emu(round(x * emu_px)), Emu(round(y * emu_px)),
                                              Emu(max(1, round(w * emu_px))), Emu(max(1, round(b["h"] * emu_px))))
                tf = tb.text_frame
                tf.word_wrap = b["lines"] > 1
                tf.auto_size = MSO_AUTO_SIZE.NONE
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                para = tf.paragraphs[0]

                def fmt_para(p):
                    p.alignment = align_map.get(b["align"], PP_ALIGN.LEFT)
                    if b.get("line_height"):
                        p.line_spacing = Pt(round(b["line_height"] * pt_px, 1))

                fmt_para(para)
                for r in runs:
                    if r.get("br"):
                        para = tf.add_paragraph()
                        fmt_para(para)
                        continue
                    chunks = r["t"].split("\n") if r.get("pre") else [r["t"]]
                    for ci, chunk in enumerate(chunks):
                        if ci:
                            para = tf.add_paragraph()
                            fmt_para(para)
                        if not chunk:
                            continue
                        run = para.add_run()
                        run.text = chunk
                        f = run.font
                        name = _map_font(r["font"], swaps)
                        fonts_used[_first_family(r["font"])] = name
                        f.name = name
                        f.size = Pt(round(r["size"] * pt_px * 2) / 2)
                        f.bold = r["weight"] >= 600
                        f.italic = bool(r["italic"])
                        f.underline = bool(r["underline"])
                        c = r["color"]
                        f.color.rgb = RGBColor(int(c[0]), int(c[1]), int(c[2]))
                        rpr = run._r.get_or_add_rPr()
                        if r.get("ls"):
                            rpr.set("spc", str(int(round(r["ls"] * pt_px * 100))))
                        if r.get("strike"):
                            rpr.set("strike", "sngStrike")
                        alpha = float(c[3]) * float(b.get("opacity") or 1)
                        if alpha < 0.99:
                            from lxml import etree
                            srgb = rpr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
                            if srgb is not None:
                                a = etree.SubElement(srgb, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha")
                                a.set("val", str(int(max(0.0, alpha) * 100000)))
                text_boxes += 1
        ni = (s.get("deck_n") or (idx + 1)) - 1
        if notes and 0 <= ni < len(notes) and notes[ni]:
            slide.notes_slide.notes_text_frame.text = str(notes[ni])
    prs.save(str(out))
    return {"fonts_used": fonts_used, "text_boxes": text_boxes}


def save_target(value: Any) -> Optional[str]:
    """Validated `save_to_project_path`: a writable project path ending in .pptx."""
    from services.design import files as dfiles
    if value in (None, ""):
        return None
    if not isinstance(value, str) or not value.lower().endswith(".pptx") or not dfiles.writable(value):
        raise ValueError("save_to_project_path must be a project path ending in .pptx (e.g. exports/deck.pptx)")
    return value


def _save_into_project(ctx, out: Path, rel: str) -> Optional[int]:
    from services.design import events, versions
    from services.design import files as dfiles
    if not dfiles.write_file(ctx.pid, rel, out.read_bytes()):
        raise ValueError(f"could not write {rel} into the project")
    rec, changed = versions.snapshot(ctx.pid, "user", prompt=f"Exported PowerPoint to {rel}")
    v = rec["v"] if rec else None
    events.publish(ctx.pid, "files", {"changed": changed or [rel], "version": v})
    return v


async def export(ctx) -> Path:
    mode = ctx.options.get("mode", "screenshots")
    save_rel = save_target(ctx.options.get("save_to_project_path"))
    if mode not in ("screenshots", "editable"):
        mode = "screenshots"
    swaps = {str(k): str(v) for k, v in (ctx.options.get("fontSwaps") or {}).items()
             if isinstance(k, str) and isinstance(v, str) and v.strip()}
    cap = await _capture_slides(ctx, mode)
    info, slides = cap["info"], cap["slides"]
    n = len(slides)

    # Validation flags
    if not cap["is_deck"]:
        ctx.flag("not_a_deck", "The page has no [data-td-slide] sections; exported the first 1920×1080 viewport "
                 "as a single slide")
    hashes = []
    for s in slides:
        hsh = hashlib.sha256(s["png"])
        if s["text"]:
            hsh.update(json.dumps([[b["x"], b["y"], [r.get("t") for r in b["runs"]]] for b in s["text"]["blocks"]],
                                  sort_keys=True).encode("utf-8"))
        hashes.append(hsh.hexdigest())
    dups = [i + 1 for i in range(1, n) if hashes[i] == hashes[i - 1]]
    if dups:
        ctx.flag("duplicate_adjacent", "Slide(s) " + ", ".join(str(d) for d in dups)
                 + " render identically to the slide before — navigation may not have moved",
                 slides=dups)
    if cap["is_deck"]:
        W, H = cap["W"], cap["H"]
        bad = [i + 1 for i, s in enumerate(slides)
               if s.get("size") and (abs(s["size"][0] - W) > 1 or abs(s["size"][1] - H) > 1)]
        if bad:
            ctx.flag("slide_size_mismatch", f"Slide box differs from the declared {W}×{H} canvas on slide(s) "
                     + ", ".join(map(str, bad[:10])), slides=bad)
        elif abs(W / H - 16 / 9) > 0.01:
            ctx.flag("slide_size_mismatch", f"Deck canvas is {W}×{H}, not 16:9 — the PowerPoint slide size was "
                     "set to match", width=W, height=H)
        notes = info.get("notes")
        if not notes:
            ctx.flag("no_speaker_notes", "The deck has no speaker notes (#td-speaker-notes)")
        elif len(notes) != n:
            ctx.flag("no_speaker_notes", f"{len(notes)} speaker-note entries for {n} slides", notes=len(notes))
        forced = [i + 1 for i, s in enumerate(slides) if s["forced"]]
        if forced:
            ctx.flag("deck_nav_forced", "The deck did not respond to td:slide navigation; slides were shown by "
                     "force for capture", slides=forced)
    if cap.get("missing_selectors"):
        ctx.flag("selector_not_found", "The slide selector matched nothing on slide(s) "
                 + ", ".join(map(str, cap["missing_selectors"][:10])) + " — the whole slide was captured instead",
                 slides=cap["missing_selectors"])
    if cap.get("js_failed"):
        ctx.flag("show_js_failed", "showJs threw on slide(s) " + ", ".join(map(str, cap["js_failed"][:10])),
                 slides=cap["js_failed"])
    if cap["errors"]:
        ctx.flag("console_errors", f"{len(cap['errors'])} console error(s) while rendering",
                 errors=cap["errors"][:5])

    ctx.progress(0.85, "Building PowerPoint")
    out = ctx.out_dir / ctx.name("pptx", "-editable" if mode == "editable" else "")
    built = await asyncio.to_thread(_build, cap, mode, swaps, out)
    if mode == "editable":
        web = sorted({src for src, dst in built["fonts_used"].items()
                      if dst.lower() not in _OFFICE_FONTS and src.lower() not in {k.lower() for k in swaps}})
        if web:
            ctx.flag("web_fonts", "PowerPoint will substitute these unless installed or swapped (fontSwaps): "
                     + ", ".join(web), fonts=web)
        ctx.job["stats"] = {"slides": n, "text_boxes": built["text_boxes"]}
    else:
        ctx.job["stats"] = {"slides": n}
    if save_rel:
        v = await asyncio.to_thread(_save_into_project, ctx, out, save_rel)
        ctx.job["stats"]["saved_to"] = save_rel
        if v:
            ctx.job["stats"]["version"] = v
    return out
