"""TeleDesign parity gaps — UI / export / share side:

share snapshots (frozen files, changes-since, update, expiry, revoke, recipient ZIP,
missing-dependency check, preview-origin route), undo/redo as version steps, chat
download route, PPTX options (googleFontImports, resetTransformSelector, per-slide
selector/showJs/delay, save_to_project_path), handoff into a Task-mode session,
Send to Google Slides via `gws`, Figma link import, the no-JS standalone fallback,
the welcome sample project and `brandFonts`.

Nothing here touches the real data/ or settings.json: design paths resolve under
tmp_path, settings reads are overridden, `config.set_nested` is captured, the task
queue's submit is stubbed (no CLI ever runs), gws and Figma are faked (no Google or
Figma call). The Edge test uses a throwaway profile and is skipped without Edge.
"""

from __future__ import annotations

import asyncio
import io
import json
import socket
import zipfile
from pathlib import Path

import pytest

import config


def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


@pytest.fixture
def td(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "_settings_dir", lambda: tmp_path)
    overrides = {
        "design.preview_port": _free_port(),
        "proxy.port": _free_port(),
        "design.share.enabled": True,
        "design.welcome_project": True,
        "design.figma.token": "",
    }
    orig = config.get_nested

    def get_nested(path, default=None):
        return overrides[path] if path in overrides else orig(path, default)
    monkeypatch.setattr(config, "get_nested", get_nested)
    written = {}

    def set_nested(path, value):          # never write the real settings.json
        written[path] = value
        overrides[path] = value
    monkeypatch.setattr(config, "set_nested", set_nested)
    from services.design import store
    return {"store": store, "overrides": overrides, "written": written, "root": tmp_path}


def _project(store, files=None, title="Gaps UI"):
    rec = store.create_project({"title": title})
    d = store.project_dir(rec["id"])
    for rel, text in (files or {}).items():
        p = d / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(text if isinstance(text, bytes) else text.encode("utf-8"))
    return rec["id"], d


def _app():
    from aiohttp import web
    from proxy import api_design
    app = web.Application()
    api_design.register_routes(app)
    return app


def _run_app(go, app=None):
    from aiohttp.test_utils import TestClient, TestServer

    async def main():
        async with TestClient(TestServer(app or _app())) as client:
            return await go(client)
    return asyncio.run(main())


# ── #9 share snapshots ───────────────────────────────────────────────────

def test_share_snapshot_is_frozen_and_tracks_changes(td):
    from services.design import share
    pid, d = _project(td["store"], {"index.html": "<h1>v1</h1>", "_ds/x/tokens.css": ":root{}"})
    rec = share.create_share(pid, "view")
    assert rec["live"] is False and rec["changed_count"] == 0 and rec["snapshot"]["file_count"] == 2
    full = share._load()[rec["token"]]
    # Edit after sharing: the recipient still sees v1; the owner sees one change.
    (d / "index.html").write_text("<h1>v2</h1>", encoding="utf-8")
    (d / "new.html").write_text("<p>new</p>", encoding="utf-8")
    assert share.read_file(full, "index.html") == b"<h1>v1</h1>"
    assert share.read_file(full, "new.html") is None
    assert share.read_file(full, "_ds/x/tokens.css") == b":root{}"      # staged DS travels with it
    [pub] = share.list_shares(pid)
    assert pub["changes"] == {"added": ["new.html"], "modified": ["index.html"], "removed": []}
    # Update in place: same token, new content.
    up = share.update_share(pid, rec["token"], resnapshot=True)
    assert up["token"] == rec["token"] and up["changed_count"] == 0
    assert share.read_file(share._load()[rec["token"]], "index.html") == b"<h1>v2</h1>"


def test_share_expiry_role_revoke_and_legacy(td, monkeypatch):
    from services.design import share
    pid, _ = _project(td["store"], {"index.html": "<p>x</p>"})
    with pytest.raises(ValueError):
        share.create_share(pid, "view", expires_in=5)            # below the 60 s floor
    rec = share.create_share(pid, "comment", expires_in=3600)
    assert rec["expires_at"] and share.resolve(rec["token"])
    data = share._load()
    data[rec["token"]]["expires_at"] = "2000-01-01T00:00:00Z"
    share._save(data)
    assert share.resolve(rec["token"]) is None                   # expired = gone
    assert share.list_shares(pid)[0]["expired"] is True
    share.update_share(pid, rec["token"], expires_in=None, role="edit")
    assert share.resolve(rec["token"])["role"] == "edit"
    # Records from before snapshots keep serving the live project.
    data = share._load()
    data["L" * 43] = {"token": "L" * 43, "project_id": pid, "role": "view", "created_at": "2026-01-01T00:00:00Z"}
    share._save(data)
    legacy = share.resolve("L" * 43)
    assert legacy and share.read_file(legacy, "index.html") == b"<p>x</p>"
    assert share.delete_share(pid, rec["token"]) and share.resolve(rec["token"]) is None
    monkeypatch.setitem(td["overrides"], "design.share.enabled", False)
    assert share.resolve("L" * 43) is None


def test_share_missing_dependencies_and_zip(td):
    from services.design import share
    pid, _ = _project(td["store"], {
        "index.html": '<img src="assets/logo.png"><link href="style.css" rel=stylesheet>'
                      '<script src="https://cdn.x/y.js"></script><a href="#top">t</a><img src="/_td/x.js">',
        "style.css": "body{background:url('img/bg.png')}",
        "sub/page.html": '<img src="../assets/logo.png"><img src="ok.png">', "sub/ok.png": b"\x89PNG"})
    rec = share.create_share(pid, "view")
    deps = share.dependency_report(share._load()[rec["token"]])
    missing = sorted({d["missing"] for d in deps})
    assert missing == ["assets/logo.png", "img/bg.png"]
    name, data = share.recipient_zip(share._load()[rec["token"]])
    z = zipfile.ZipFile(io.BytesIO(data))
    assert name == "gaps-ui.zip"
    assert sorted(n.split("/", 1)[1] for n in z.namelist()) == ["index.html", "style.css", "sub/ok.png", "sub/page.html"]


def test_share_routes_and_preview_snapshot(td):
    from aiohttp import web
    from services.design import preview
    pid, d = _project(td["store"], {"index.html": "<html><head></head><body><h1>Snap</h1></body></html>",
                                    "comments.json": "[]"})

    async def go(c):
        r = await c.post(f"/api/design/projects/{pid}/share", json={"role": "view", "expires_in": 86400})
        assert r.status == 200
        body = await r.json()
        tok = body["share"]["token"]
        assert body["missing_dependencies"] == []
        (d / "index.html").write_text("<h1>changed</h1>", encoding="utf-8")
        lst = await (await c.get(f"/api/design/projects/{pid}/share")).json()
        assert lst["shares"][0]["changed_count"] == 1
        meta = await (await c.get(f"/api/design/s/{tok}")).json()
        assert meta["pages_base"] == f"/s/{tok}/" and [f["path"] for f in meta["files"]] == ["index.html"]
        assert "design_system_id" not in meta["project"]
        f = await c.get(f"/api/design/s/{tok}/files/index.html")
        assert "Snap" in await f.text() and f.headers["Content-Type"].startswith("text/plain")
        z = await c.get(f"/api/design/s/{tok}/download")
        assert z.status == 200 and z.headers["Content-Type"] == "application/zip"
        chk = await (await c.get(f"/api/design/projects/{pid}/share/{tok}/check")).json()
        assert chk["share"]["changed_count"] == 1
        up = await c.patch(f"/api/design/projects/{pid}/share/{tok}", json={"resnapshot": True})
        assert (await up.json())["share"]["changed_count"] == 0
        # A foreign Origin can't write (the preview origin included).
        bad = await c.patch(f"/api/design/projects/{pid}/share/{tok}", json={"resnapshot": True},
                            headers={"Origin": "http://127.0.0.1:1"})
        assert bad.status == 403
        return tok

    tok = _run_app(go)

    async def prev(c):
        r = await c.get(f"/s/{tok}/index.html")
        text = await r.text()
        assert r.status == 200 and "td-bridge.js" in text and "changed" in text
        assert (await c.get(f"/s/{tok}/comments.json")).status == 404
        assert (await c.get(f"/s/{'x' * 43}/index.html")).status == 404
        from services.design import share
        share.delete_share(pid, tok)
        assert (await c.get(f"/s/{tok}/index.html")).status == 404
    _run_app(prev, app=preview.create_app())


# ── #15 undo / redo as version steps ─────────────────────────────────────

def test_undo_redo_steps_are_versions(td):
    from services.design import versions
    pid, d = _project(td["store"])
    for text in ("A", "B", "C"):
        (d / "page.html").write_text(text, encoding="utf-8")
        versions.snapshot(pid, "user")
    n0 = len(versions.list_versions(pid))
    assert versions.undo_state(pid, "page.html") == {"can_undo": True, "can_redo": False, "pos": 2, "count": 3}
    r = versions.step(pid, "page.html", "undo")
    assert (d / "page.html").read_text() == "B" and r["to_v"] == 2 and r["can_redo"]
    assert r["version"]["origin"] == "restore" and r["version"]["undo"]["direction"] == "undo"
    versions.step(pid, "page.html", "undo")
    assert (d / "page.html").read_text() == "A"
    assert versions.step(pid, "page.html", "undo") is None          # nothing earlier
    versions.step(pid, "page.html", "redo")
    assert (d / "page.html").read_text() == "B"
    assert len(versions.list_versions(pid)) == n0 + 3               # non-destructive: every step recorded
    # A fresh edit drops the redo branch and becomes the head.
    (d / "page.html").write_text("D", encoding="utf-8")
    st = versions.undo_state(pid, "page.html")
    assert st["can_redo"] is False and st["can_undo"] is True
    versions.step(pid, "page.html", "undo")                          # snapshots D first, then goes back
    assert (d / "page.html").read_text() == "B"
    assert any(v.get("prompt", "").startswith("Before undo") for v in versions.list_versions(pid))
    with pytest.raises(ValueError):
        versions.step(pid, "../x", "undo")


def test_undo_route(td):
    from services.design import versions
    pid, d = _project(td["store"])
    for text in ("one", "two"):
        (d / "a.html").write_text(text, encoding="utf-8")
        versions.snapshot(pid, "user")

    async def go(c):
        s = await (await c.get(f"/api/design/projects/{pid}/versions/undo?path=a.html")).json()
        assert s["can_undo"] is True
        r = await (await c.post(f"/api/design/projects/{pid}/versions/undo", json={"path": "a.html", "direction": "undo"})).json()
        assert r["ok"] and r["to_v"] == 1
        r2 = await (await c.post(f"/api/design/projects/{pid}/versions/undo", json={"path": "a.html"})).json()
        assert r2["ok"] is False
        bad = await c.post(f"/api/design/projects/{pid}/versions/undo", json={"path": ".versions/manifest.json"})
        assert bad.status == 400
    _run_app(go)
    assert (d / "a.html").read_text() == "one"


def test_bridge_forwards_undo_keys():
    src = (Path(__file__).resolve().parent.parent / "services/design/runtime/td-bridge.js").read_text(encoding="utf-8")
    assert "type: 'td:key'" in src
    host = (Path(__file__).resolve().parent.parent / "proxy/static/design/app/bridge.js").read_text(encoding="utf-8")
    assert '"td:key"' in host


# ── #8 download cards: the download route ────────────────────────────────

def test_download_route(td):
    pid, _ = _project(td["store"], {"index.html": "<p>hi</p>", "exports/a.txt": "A", "exports/b/c.txt": "C",
                                    "comments.json": "[]"})

    async def go(c):
        P = f"/api/design/projects/{pid}/download"
        f = await c.get(P + "?path=index.html")
        assert f.status == 200 and "attachment" in f.headers["Content-Disposition"]
        assert f.headers["Content-Type"] == "application/octet-stream"
        z = await c.get(P + "?path=exports/&kind=folder")
        names = zipfile.ZipFile(io.BytesIO(await z.read())).namelist()
        assert sorted(n.split("/", 1)[1] for n in names) == ["a.txt", "b/c.txt"]
        # A folder path without kind= still zips (file lookup falls through).
        assert (await c.get(P + "?path=exports")).headers["Content-Type"] == "application/zip"
        proj = zipfile.ZipFile(io.BytesIO(await (await c.get(P + "?kind=project")).read())).namelist()
        assert not any(n.endswith("comments.json") for n in proj) and any(n.endswith("index.html") for n in proj)
        for bad in ("../x", ".versions/manifest.json", "nope.txt", "chats/"):
            assert (await c.get(P + "?path=" + bad)).status in (400, 404)
    _run_app(go)


# ── #11 PPTX options ─────────────────────────────────────────────────────

def test_pptx_option_parsing():
    from services.design import pptx_export as px
    urls = px.font_import_urls(["Inter:wght@400;700", "Newsreader", "https://fonts.googleapis.com/css2?family=Lora",
                                "https://evil.example/x.css", "bad'); x", 12])
    assert urls == ["https://fonts.googleapis.com/css2?family=Inter:wght@400;700&display=swap",
                    "https://fonts.googleapis.com/css2?family=Newsreader&display=swap",
                    "https://fonts.googleapis.com/css2?family=Lora"]
    specs = px.slide_specs({"slides": [{"index": 2, "selector": "#a", "showJs": "go()", "delay": 99999},
                                       {"selector": "a{b}", "delay": "x"}, "junk"]})
    assert specs[0] == {"index": 2, "selector": "#a", "showJs": "go()", "delay": 10000}
    assert specs[1] == {"index": None, "selector": None, "showJs": None, "delay": 0}
    assert px.slide_specs({}) is None and px.slide_specs({"slides": []}) is None
    assert px.save_target("exports/deck.pptx") == "exports/deck.pptx" and px.save_target(None) is None
    for bad in ("../deck.pptx", "deck.pdf", "_ds/x.pptx", "C:/x.pptx"):
        with pytest.raises(ValueError):
            px.save_target(bad)


def _edge() -> bool:
    try:
        from services.design import render
        render._edge_binary()
        import pptx  # noqa: F401
        return True
    except Exception:
        return False


needs_edge = pytest.mark.skipif(not _edge(), reason="Edge or python-pptx not installed")

_DECK = """<!doctype html><html><head><meta charset="utf-8"><style>
body{margin:0} #stage{transform:scale(.5);transform-origin:0 0}
section{width:1920px;height:1080px;display:none;font:64px sans-serif}
section.on{display:block} .tab{display:none} .tab.on{display:block}</style></head><body>
<div id="stage" data-td-deck data-td-w="1920" data-td-h="1080">
<section data-td-slide class="on"><h1>One</h1><div id="card" style="width:800px;height:400px;background:#36c">Card</div></section>
<section data-td-slide><h1>Two</h1><div class="tab" id="t2">Hidden tab</div></section>
</div>
<script type="application/json" id="td-speaker-notes">["n1", "n2"]</script>
<script>const s=[...document.querySelectorAll('section')];addEventListener('message',e=>{const d=e.data||{};
if(d.type==='td:slide'&&d.action==='go'){s.forEach((x,k)=>x.classList.toggle('on',k===d.index-1))}});</script>
</body></html>"""


@needs_edge
def test_pptx_export_with_options_in_edge(td):
    from pptx import Presentation
    from services.design import export, preview, render, versions
    pid, d = _project(td["store"], {"deck.html": _DECK})

    async def go():
        from aiohttp import web
        runner = web.AppRunner(preview.create_app())
        await runner.setup()
        await web.TCPSite(runner, "127.0.0.1", td["overrides"]["design.preview_port"]).start()
        try:
            return await export.run_sync(pid, "pptx", {"file": "deck.html", "options": {
                "mode": "editable", "resetTransformSelector": "#stage",
                "googleFontImports": ["Inter"],
                "slides": [{"index": 1, "selector": "#card"},
                           {"index": 2, "showJs": "document.getElementById('t2').classList.add('on')", "delay": 50},
                           {"index": 2, "selector": "#nope"}],
                "save_to_project_path": "exports/deck.pptx"}})
        finally:
            await runner.cleanup()
            render.stop()
    job = asyncio.run(go())
    assert job["status"] == "done", job.get("error")
    assert job["stats"]["saved_to"] == "exports/deck.pptx" and (d / "exports/deck.pptx").is_file()
    codes = {f["code"] for f in job["flags"]}
    assert "selector_not_found" in codes and "show_js_failed" not in codes
    prs = Presentation(str(d / "exports/deck.pptx"))
    slides = list(prs.slides)
    assert len(slides) == 3
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in slides]
    assert "Card" in texts[0] and "One" not in texts[0]          # selector region only
    assert "Hidden tab" in texts[1]                               # showJs ran before capture
    assert [s.notes_slide.notes_text_frame.text for s in slides] == ["n1", "n2", "n2"]
    assert any(v.get("prompt", "").startswith("Exported PowerPoint") for v in versions.list_versions(pid))


# ── lower-priority: no-JS fallback in standalone HTML ────────────────────

def test_standalone_noscript_fallback(td, monkeypatch):
    from PIL import Image
    from services.design import export, render
    pid, _ = _project(td["store"], {"index.html": "<!doctype html><html><head></head><body><div id=root></div></body></html>"})
    buf = io.BytesIO()
    Image.new("RGB", (1280, 2000), "#abc").save(buf, "JPEG")

    async def shot(*a, **kw):
        assert kw.get("full_page") is True
        return buf.getvalue()
    monkeypatch.setattr(render, "screenshot", shot)
    doc, _ = asyncio.run(export.build_standalone(pid, "index.html"))
    assert "<noscript><style>body>*:not(noscript){display:none !important}</style>" in doc
    assert 'src="data:image/webp;base64,' in doc.split("<noscript>", 1)[1]
    assert '<template id="__bundler_thumbnail">' in doc
    doc2, _ = asyncio.run(export.build_standalone(pid, "index.html", splash=False, noscript=False))
    assert "<noscript>" not in doc2 and "__bundler_thumbnail" not in doc2
    assert "needs JavaScript" in export._noscript_fallback(None, "T")


# ── #13 handoff into a Task-mode session ─────────────────────────────────

def test_handoff_list_dirs(td, tmp_path):
    from services.design import handoff
    (tmp_path / "repo" / ".git").mkdir(parents=True)
    (tmp_path / "plain").mkdir()
    (tmp_path / ".hidden").mkdir()
    (tmp_path / "file.txt").write_text("x")
    r = handoff.list_dirs(str(tmp_path))
    names = {x["name"]: x["git"] for x in r["dirs"]}
    assert names.get("repo") is True and names.get("plain") is False and ".hidden" not in names and "file.txt" not in names
    with pytest.raises(ValueError):
        handoff.list_dirs("relative/path")
    with pytest.raises(ValueError):
        handoff.list_dirs(str(tmp_path / "file.txt"))


def test_handoff_session_stages_bundle_and_submits(td, tmp_path, monkeypatch):
    from services.design import handoff
    from services.session import session_store
    from services.task import task_manager
    monkeypatch.setattr(session_store, "get_sessions_dir", lambda: tmp_path / "sessions")
    calls = []

    class Q:
        def submit_task(self, **kw):
            calls.append(kw)
            return "task-123"
    monkeypatch.setattr(task_manager, "get_task_queue", lambda: Q())
    pid, _ = _project(td["store"], {"index.html": "<h1>Design</h1>"})
    repo = tmp_path / "my-app"
    repo.mkdir()
    res = handoff.start_session(pid, {"repo": str(repo), "engine": "codex", "model": "gpt-5", "note": "Only the hero."})
    assert res["task_id"] == "task-123" and res["task_type"] == "CODEX"
    [call] = calls
    assert call["task_type"] == "CODEX" and call["session_id"] == res["session_id"]
    assert call["params"]["model"] == "gpt-5" and call["params"]["is_local"] is False
    prompt = call["params"]["prompt"]
    assert str(repo.resolve()) in prompt and "handoff/README.md" in prompt and "Only the hero." in prompt
    folder = session_store._session_dir(res["session_id"], None)
    assert (folder / "handoff" / "README.md").is_file() and (folder / "handoff" / "project" / "index.html").is_file()
    meta = session_store.get(res["session_id"])
    assert meta["data"]["repo"] == str(repo.resolve()) and meta["session_idle_timeout_seconds"] == 0
    for bad in ({"repo": "rel"}, {"repo": str(tmp_path / "missing")}, {"repo": str(repo), "engine": "vim"},
                {"repo": str(repo), "model": "a b"}):
        with pytest.raises(ValueError):
            handoff.start_session(pid, bad)
    assert len(calls) == 1


# ── #19 Send to Google Slides via gws ────────────────────────────────────

class _CP:
    def __init__(self, out=b"", err=b"", code=0):
        self.stdout, self.stderr, self.returncode = out, err, code


def test_gslides_status_never_logs_in(td, monkeypatch):
    from services.design import gslides
    monkeypatch.setattr(gslides, "binary", lambda: None)
    st = gslides.status(refresh=True)
    assert st["installed"] is False and "gws auth login" in st["instructions"]
    monkeypatch.setattr(gslides, "binary", lambda: "C:/fake/gws.exe")
    seen = []

    def run(args, timeout=60.0):
        seen.append(args)
        return _CP(b'Using keyring backend: keyring\n{"token_valid": false, "has_refresh_token": false}')
    monkeypatch.setattr(gslides, "_run", run)
    st = gslides.status(refresh=True)
    assert st["installed"] and not st["authed"] and "gws auth login" in st["instructions"]
    assert seen == [["auth", "status"]]                               # read-only check, no login
    monkeypatch.setattr(gslides, "_run", lambda args, timeout=60.0: _CP(
        b'{"token_valid": true, "user": "a@b.c", "scopes": ["https://www.googleapis.com/auth/drive"]}'))
    st = gslides.status(refresh=True)
    assert st["authed"] and st["can_upload"] and st["user"] == "a@b.c"


def test_gslides_upload_and_job(td, monkeypatch):
    from services.design import export, gslides
    pid, d = _project(td["store"], {"index.html": "<h1>x</h1>"})
    args_seen = []

    def run(args, timeout=60.0):
        args_seen.append(args)
        return _CP(b'{"id": "F1", "name": "Gaps UI", "webViewLink": "https://docs.google.com/presentation/d/F1/edit"}')
    monkeypatch.setattr(gslides, "_run", run)
    out = gslides.upload(d / "index.html", "Gaps UI")
    assert out == {"id": "F1", "url": "https://docs.google.com/presentation/d/F1/edit", "name": "Gaps UI"}
    a = args_seen[0]
    assert a[:3] == ["drive", "files", "create"] and json.loads(a[a.index("--json") + 1])["mimeType"] == gslides.SLIDES_MIME
    monkeypatch.setattr(gslides, "_run", lambda args, timeout=60.0: _CP(b"", b"403 insufficient scope", 1))
    with pytest.raises(RuntimeError, match="insufficient scope"):
        gslides.upload(d / "index.html", "x")

    monkeypatch.setattr(gslides, "status", lambda refresh=False: {"installed": True, "authed": True, "can_upload": True})
    fake_pptx = d / "deck.pptx"
    fake_pptx.write_bytes(b"PK")

    def fake_start(pid_, kind, body):
        assert kind == "pptx" and body["options"]["mode"] == "editable" and "save_to_project_path" not in body["options"]
        return {"id": "j", "status": "done", "progress": 1.0, "result": "deck.pptx"}
    monkeypatch.setattr(export, "start", fake_start)
    monkeypatch.setattr(export, "result_path", lambda job: fake_pptx)
    monkeypatch.setattr(gslides, "upload", lambda path, name: {"id": "F2", "url": "https://docs.google.com/presentation/d/F2/edit", "name": name})

    async def go():
        job = gslides.start(pid, {"file": "index.html", "options": {"save_to_project_path": "x.pptx"}})
        for _ in range(50):
            await asyncio.sleep(0.05)
            j = gslides.get_job(job["id"])
            if j["status"] in ("done", "failed"):
                return j
    j = asyncio.run(go())
    assert j["status"] == "done" and j["url"].endswith("/F2/edit")


def test_gslides_route_refuses_when_signed_out(td, monkeypatch):
    from services.design import gslides
    pid, _ = _project(td["store"], {"index.html": "<h1>x</h1>"})
    monkeypatch.setattr(gslides, "status", lambda refresh=False: {"installed": True, "authed": False, "instructions": gslides.LOGIN_HINT})

    async def go(c):
        r = await c.post(f"/api/design/projects/{pid}/send/google-slides", json={"file": "index.html"})
        assert r.status == 412 and "gws auth login" in (await r.json())["error"]
        s = await (await c.get("/api/design/gslides/status")).json()
        assert s["authed"] is False
    _run_app(go)


# ── Figma link import ────────────────────────────────────────────────────

def test_figma_link_parsing_and_token(td):
    from services.design import figma_import as fi
    assert fi.parse_link("https://www.figma.com/design/AbCdEf12345/My-File?node-id=12-34") == {"key": "AbCdEf12345", "node_id": "12:34"}
    assert fi.parse_link("https://figma.com/file/AbCdEf12345/x")["node_id"] is None
    for bad in ("http://www.figma.com/design/AbCdEf12345/x", "https://evil.com/design/AbCdEf12345", "https://www.figma.com/about"):
        with pytest.raises(ValueError):
            fi.parse_link(bad)
    with pytest.raises(ValueError):
        fi.set_token("not a token!")
    fi.set_token("figd_" + "a" * 30)
    assert td["written"] == {"design.figma.token": "figd_" + "a" * 30} and fi.token()


def test_figma_import_frames(td, monkeypatch):
    from services.design import figma_import as fi
    pid, d = _project(td["store"])
    calls = []

    async def get_json(path, params=None):
        calls.append((path, params))
        if path == "/files/AbCdEf12345":
            return {"name": "Mobile", "document": {"children": [{"name": "Page 1", "children": [
                {"id": "1:2", "type": "FRAME", "name": "Home screen", "absoluteBoundingBox": {"width": 390, "height": 844}},
                {"id": "1:3", "type": "TEXT", "name": "loose text"},
                {"id": "1:4", "type": "FRAME", "name": "Home screen", "absoluteBoundingBox": {"width": 390, "height": 900}}]}]}}
        if path == "/images/AbCdEf12345":
            assert params["ids"] == "1:2,1:4" and params["scale"] == "2"
            return {"images": {"1:2": "https://s3.example/a.png", "1:4": "https://s3.example/b.png"}}
        raise AssertionError(path)

    async def fetch(url):
        assert url.startswith("https://s3.example/")
        return b"\x89PNG-" + url[-5:].encode()
    monkeypatch.setattr(fi, "_get_json", get_json)
    monkeypatch.setattr(fi, "_fetch_bytes", fetch)
    listing = asyncio.run(fi.list_frames("https://www.figma.com/design/AbCdEf12345/x"))
    assert [f["id"] for f in listing["frames"]] == ["1:2", "1:4"]
    res = asyncio.run(fi.import_frames(pid, "https://www.figma.com/design/AbCdEf12345/x"))
    assert res["boards"] == ["figma-home-screen.html", "figma-home-screen-2.html"]
    assert (d / "imports/figma/home-screen.png").read_bytes().startswith(b"\x89PNG")
    page = (d / "figma-home-screen.html").read_text(encoding="utf-8")
    assert 'src="imports/figma/home-screen.png"' in page and 'width="390"' in page
    with pytest.raises(fi.FigmaError):
        asyncio.run(fi.import_frames(pid, "https://www.figma.com/design/AbCdEf12345/x", ids=["9:9"]))


def test_figma_needs_a_token(td):
    from services.design import figma_import as fi
    with pytest.raises(fi.FigmaError, match="token"):
        asyncio.run(fi._get_json("/files/x"))


# ── Welcome sample project + brandFonts ──────────────────────────────────

def test_welcome_project_once(td, monkeypatch):
    from services.design import templates
    store = td["store"]
    p = templates.ensure_welcome()
    assert p and p["title"] == "Welcome to TeleDesign" and p.get("sample") is True
    assert (store.project_dir(p["id"]) / "index.html").is_file()
    assert templates.ensure_welcome() is None                           # marker: never twice
    store.delete_project(p["id"])
    assert templates.ensure_welcome() is None and store.list_projects() == []


def test_welcome_skipped_when_projects_exist_or_disabled(td, monkeypatch):
    from services.design import templates
    _project(td["store"])
    assert templates.ensure_welcome() is None
    monkeypatch.setitem(td["overrides"], "design.welcome_project", False)
    (templates._welcome_marker()).unlink()
    assert templates.ensure_welcome() is None


def test_welcome_route(td):
    async def go(c):
        r = await (await c.post("/api/design/welcome", json={})).json()
        assert r["project"]["title"] == "Welcome to TeleDesign"
        r2 = await (await c.post("/api/design/welcome", json={})).json()
        assert r2["project"] is None
    _run_app(go)


def test_brand_fonts_normalised():
    from services.design import systems
    assert systems.brand_fonts({"brandFonts": [{"family": "Söhne", "status": "substituted", "substitute": "Inter",
                                                "tokens": ["--font-sans"]}]}) == [
        {"family": "Söhne", "status": "substituted", "tokens": ["--font-sans"], "substitute": "Inter"}]
    cd = systems.brand_fonts({"brandFonts": {"status": "provided", "tokens": {"--font-display": "'Newsreader', serif",
                                                                               "--font-body": "Inter"}}})
    assert {x["family"] for x in cd} == {"Newsreader", "Inter"} and all(x["status"] == "provided" for x in cd)
    assert systems.brand_fonts({"brandFonts": {"Brand Grotesk": "missing"}})[0]["status"] == "missing"
    assert systems.brand_fonts({}) == [] and systems.brand_fonts({"brandFonts": "junk"}) == []
    man = json.loads(systems.compact_manifest({"name": "X", "fonts": [{"family": "Inter", "style": "italic", "remoteSrc": "https://f/x.woff2"}],
                                               "brandFonts": [{"family": "Inter", "status": "provided"}]}))
    assert man["brandFonts"][0]["family"] == "Inter" and man["fonts"][0]["style"] == "italic"
    assert man["fonts"][0]["remoteSrc"] == "https://f/x.woff2"
